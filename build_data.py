"""Lee banco_problemas.xlsx y genera data.js (modelos de barras incluidos).
Uso:  python build_data.py   (necesita: pip install openpyxl)"""
import re, json, sys
from openpyxl import load_workbook
src = sys.argv[1] if len(sys.argv) > 1 else "banco_problemas.xlsx"
ws = load_workbook(src, data_only=True)["Problemas"]
G = {"v": None}
out = []
for r in ws.iter_rows(min_row=2, values_only=True):
    id_, mundo, lvl, edad, tipo, tx, a, b, c, resp, op, modelo, h1, h2, err = r[:15]
    w, t = int(id_[1]), (int(id_[-2:]) - 1) // 3
    m = re.match(r"(\S+) tiene .* y (\S+) tiene", tx); n1, n2 = m.groups() if m else (tx.split()[0], None)
    tot = None; rows = None; fr = None; nn = None
    if w == 1:
        rows = [{"l": "Caja" if t == 2 else n1, "s": [a, b] if t == 0 else [b, "?"]}]
        tot, val = ("?", a + b) if t == 0 else (a, a - b)
    elif w == 2:
        if t == 0: rows = [{"l": n1, "s": [a]}, {"l": n2, "s": [a, b], "t": "?"}]; val = a + b
        elif t == 1: rows = [{"l": n1, "s": [a]}, {"l": n2, "s": ["?", {"v": b, "g": 1}]}]; val = a - b
        elif t == 2: rows = [{"l": n1, "s": [b, "?"], "t": a}, {"l": n2, "s": [b]}]; val = a - b
        else: rows = [{"l": n1, "s": [a]}, {"l": n2, "s": [a, b]}]; tot, val = "?", 2 * a + b
    elif w == 3:
        d_, k_ = b, c
        rows = [{"l": "Todo", "s": []}]
        if t == 0: fr, tot, val, nn = {"d": d_, "k": k_, "g1": "?"}, a, a // d_ * k_, 5
        elif t == 1: fr, tot, val, nn = {"d": d_, "k": k_, "g2": "?"}, a, a - a // d_ * k_, 7
        elif t == 2: fr, tot, val, nn = {"d": d_, "k": 1, "g2": "?"}, a, a - a // d_, 5
        else: fr, tot, val, nn = {"d": d_, "k": k_, "g1": a}, "?", a // k_ * d_, 5
    elif w == 4 and b <= 6 and t < 3:
        if t == 0: rows = [{"l": "Grupos", "s": [a] * b}]; tot, val = "?", a * b
        elif t == 1: rows = [{"l": "Cajas", "s": ["?"] * b}]; tot, val = a, a // b
        else: rows = [{"l": re.search(r"de (\w+)\.", tx).group(1), "s": [a]}, {"l": n1, "s": [a] * b}]; tot, val = "?", a * (b + 1)
    if rows is None: continue
    if resp is not None: assert resp == val, (id_, resp, val)
    nums = {x for rw in ([] if fr else rows) for s in rw["s"] for x in [s["v"] if isinstance(s, dict) else s] if x != "?"}
    nums |= {x for x in [tot] + [rw.get("t") for rw in rows] if isinstance(x, int)}
    if fr: nums |= {x for x in [fr.get("g1"), fr.get("g2")] if isinstance(x, int)}
    d = val + 1 if val + 1 not in nums else val + 2
    p = {"id": id_, "w": w, "l": lvl, "t": tx, "rows": rows, "ans": val, "tray": sorted(nums) + ["?", d],
         "ops": sorted({a, b} | ({c} if w == 3 else set())) + ["+", "−", "×", "÷"], "n": nn or (5 if (w == 2 and t == 3) or (w == 4 and t == 2) else 3),
         "h": [h1, h2, "Prueba con: " + op.replace("(", "").replace(")", "")]}
    if tot is not None: p["tot"] = tot
    if fr: p["fr"] = fr
    out.append(p)
open("data.js", "w", encoding="utf-8").write("const DATA=" + json.dumps({"problems": out}, ensure_ascii=False) + ";\n")
print(len(out), "problemas con modelo de barras")
